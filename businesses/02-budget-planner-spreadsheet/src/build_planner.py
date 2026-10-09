"""Generate the "Smart Budget Planner" spreadsheet in English and Arabic (Excel + Google Sheets compatible).

Usage: python3 build_planner.py <output_dir>
"""
import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
TEAL, TEAL_LIGHT, YELLOW, GREY = "0F766E", "CCFBF1", "FFF59D", "F3F4F6"
RED_TXT, GREEN_TXT = "B91C1C", "15803D"
ROWS = 1000  # transaction rows

L = {
    "en": dict(
        file="Smart-Budget-Planner-EN.xlsx", rtl=False,
        income="Income", expense="Expense",
        sheets=["Start Here", "Settings", "Transactions", "Dashboard", "Savings Goals", "Debt Payoff", "Bills"],
        title="Smart Budget Planner", subtitle="Track every dollar · see where your money goes · hit your goals",
        steps=[
            "1. Go to Settings: set your currency and edit the categories + monthly budget (yellow cells).",
            "2. Go to Transactions: log each income or expense. Pick Type and Category from the dropdowns.",
            "3. Open Dashboard: pick a month at the top to see budget vs actual, charts and savings rate.",
            "4. Savings Goals: add goals; see the % done and how much to save each month.",
            "5. Debt Payoff: enter each debt; see months to debt-free and total interest.",
            "6. Bills: list recurring bills and tick each month when paid.",
        ],
        legend="Legend: YELLOW cells = type here. White cells = automatic formulas (don't edit).",
        tip="Google Sheets: File → Import → Upload → 'Replace spreadsheet'. Excel: just open it.",
        example_note="Example rows are included so you can see how it works. Delete them and start your own.",
        settings_hdr=["Category", "Type", "Monthly Budget", "Spent This Year", "Avg / Month"],
        currency_lbl="Currency symbol", year_lbl="Year", start_lbl="Starting balance",
        cats=[("Salary", "I", 0), ("Side Hustle", "I", 0), ("Gifts / Other", "I", 0),
              ("Rent / Mortgage", "E", 1200), ("Groceries", "E", 450), ("Utilities", "E", 180),
              ("Transport", "E", 200), ("Phone & Internet", "E", 80), ("Eating Out", "E", 150),
              ("Shopping", "E", 120), ("Health", "E", 60), ("Kids / Family", "E", 150),
              ("Subscriptions", "E", 40), ("Education", "E", 50), ("Charity / Zakat", "E", 50),
              ("Entertainment", "E", 60), ("Savings Transfer", "E", 300), ("Debt Payments", "E", 200),
              ("Miscellaneous", "E", 50), ("", "", None)],
        tx_hdr=["Date", "Type", "Category", "Description", "Amount", "Month"],
        examples=[(1, "I", "Salary", "Monthly salary", 3500), (2, "E", "Rent / Mortgage", "Rent", 1200),
                  (3, "E", "Groceries", "Supermarket", 96.4), (5, "E", "Transport", "Fuel", 55),
                  (8, "E", "Eating Out", "Pizza night", 32.5), (12, "E", "Utilities", "Electricity", 74),
                  (15, "I", "Side Hustle", "Freelance project", 400), (18, "E", "Subscriptions", "Streaming", 15.99),
                  (20, "E", "Savings Transfer", "Emergency fund", 300), (24, "E", "Groceries", "Weekly shop", 88.2)],
        months=["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"],
        dash_title="Dashboard", month_pick="Choose month →",
        kpis=["Income", "Expenses", "Net (saved)", "Savings rate"],
        bva_hdr=["Category", "Budget", "Actual", "Difference", "Status"],
        ok="✓ On track", over="✗ Over budget",
        year_hdr=["Month", "Income", "Expenses", "Net", "Savings rate", "Balance"],
        chart_year="Income vs Expenses", chart_pie="Where the money went (selected month)",
        goals_hdr=["Goal", "Target", "Saved so far", "Progress", "Remaining", "Target date", "Save / month"],
        goals_ex=[("Emergency fund", 6000, 1500, 12), ("Holiday", 2000, 400, 8), ("New laptop", 1200, 300, 5)],
        debt_hdr=["Debt", "Balance", "Interest % / yr", "Monthly payment", "Months to pay off", "Debt-free by", "Total interest"],
        debt_ex=[("Credit card", 2500, 0.22, 150), ("Car loan", 9000, 0.07, 250)],
        debt_note="Tip (snowball): pay minimums on all, put every extra dollar on the smallest balance first.",
        bills_hdr=["Bill", "Due day", "Amount"],
        bills_ex=[("Rent", 1, 1200), ("Electricity", 10, 75), ("Internet", 15, 45), ("Phone", 20, 35)],
        paid_mark="✓",
    ),
    "ar": dict(
        file="Smart-Budget-Planner-AR.xlsx", rtl=True,
        income="دخل", expense="مصروف",
        sheets=["ابدأ هنا", "الإعدادات", "المعاملات", "لوحة التحكم", "أهداف الادخار", "سداد الديون", "الفواتير"],
        title="مخطط الميزانية الذكي", subtitle="تتبّع كل قرش · اعرف وين عم تروح مصاريك · حقق أهدافك",
        steps=[
            "١. افتح الإعدادات: اختر العملة وعدّل الفئات والميزانية الشهرية (الخلايا الصفراء).",
            "٢. افتح المعاملات: سجّل كل دخل أو مصروف واختر النوع والفئة من القائمة.",
            "٣. افتح لوحة التحكم: اختر الشهر لتشوف الميزانية مقابل الفعلي والرسوم ونسبة الادخار.",
            "٤. أهداف الادخار: أضف أهدافك وشوف نسبة الإنجاز وكم لازم توفّر كل شهر.",
            "٥. سداد الديون: أدخل ديونك وشوف كم شهر لتصير بلا ديون ومجموع الفوائد.",
            "٦. الفواتير: سجّل فواتيرك الشهرية وعلّم كل شهر لما تدفع.",
        ],
        legend="دليل الألوان: الخلايا الصفراء = اكتب هنا. الخلايا البيضاء = معادلات تلقائية (لا تعدّلها).",
        tip="Google Sheets: ملف ← استيراد ← تحميل. Excel: افتح الملف مباشرة.",
        example_note="في أمثلة جاهزة لتشوف كيف بيشتغل. امسحها وابدأ بأرقامك.",
        settings_hdr=["الفئة", "النوع", "الميزانية الشهرية", "المصروف هذه السنة", "المعدل الشهري"],
        currency_lbl="رمز العملة", year_lbl="السنة", start_lbl="الرصيد الافتتاحي",
        cats=[("راتب", "I", 0), ("عمل إضافي", "I", 0), ("هدايا / أخرى", "I", 0),
              ("إيجار / قرض سكن", "E", 1200), ("بقالة", "E", 450), ("كهرباء وماء", "E", 180),
              ("مواصلات", "E", 200), ("هاتف وإنترنت", "E", 80), ("مطاعم", "E", 150),
              ("تسوق", "E", 120), ("صحة", "E", 60), ("أطفال / عائلة", "E", 150),
              ("اشتراكات", "E", 40), ("تعليم", "E", 50), ("صدقة / زكاة", "E", 50),
              ("ترفيه", "E", 60), ("تحويل للادخار", "E", 300), ("أقساط ديون", "E", 200),
              ("متفرقات", "E", 50), ("", "", None)],
        tx_hdr=["التاريخ", "النوع", "الفئة", "الوصف", "المبلغ", "الشهر"],
        examples=[(1, "I", "راتب", "الراتب الشهري", 3500), (2, "E", "إيجار / قرض سكن", "الإيجار", 1200),
                  (3, "E", "بقالة", "سوبرماركت", 96.4), (5, "E", "مواصلات", "بنزين", 55),
                  (8, "E", "مطاعم", "عشاء", 32.5), (12, "E", "كهرباء وماء", "فاتورة الكهرباء", 74),
                  (15, "I", "عمل إضافي", "مشروع فريلانس", 400), (18, "E", "اشتراكات", "اشتراك شهري", 15.99),
                  (20, "E", "تحويل للادخار", "صندوق الطوارئ", 300), (24, "E", "بقالة", "تسوق أسبوعي", 88.2)],
        months=["كانون الثاني", "شباط", "آذار", "نيسان", "أيار", "حزيران", "تموز", "آب", "أيلول",
                "تشرين الأول", "تشرين الثاني", "كانون الأول"],
        dash_title="لوحة التحكم", month_pick="اختر الشهر ←",
        kpis=["الدخل", "المصاريف", "الصافي (الموفَّر)", "نسبة الادخار"],
        bva_hdr=["الفئة", "الميزانية", "الفعلي", "الفرق", "الحالة"],
        ok="✓ ضمن الميزانية", over="✗ تجاوز",
        year_hdr=["الشهر", "الدخل", "المصاريف", "الصافي", "نسبة الادخار", "الرصيد"],
        chart_year="الدخل مقابل المصاريف", chart_pie="وين راحت المصاري (الشهر المختار)",
        goals_hdr=["الهدف", "المبلغ المطلوب", "الموفَّر حتى الآن", "نسبة الإنجاز", "المتبقي", "التاريخ المستهدف", "ادخار شهري"],
        goals_ex=[("صندوق طوارئ", 6000, 1500, 12), ("سفرة", 2000, 400, 8), ("لابتوب جديد", 1200, 300, 5)],
        debt_hdr=["الدين", "الرصيد", "الفائدة السنوية %", "القسط الشهري", "أشهر حتى السداد", "تاريخ الخلاص", "مجموع الفوائد"],
        debt_ex=[("بطاقة ائتمان", 2500, 0.22, 150), ("قرض سيارة", 9000, 0.07, 250)],
        debt_note="نصيحة (كرة الثلج): ادفع الحد الأدنى للكل، وكل مبلغ زيادة حطّه على أصغر دين أولاً.",
        bills_hdr=["الفاتورة", "يوم الاستحقاق", "المبلغ"],
        bills_ex=[("الإيجار", 1, 1200), ("الكهرباء", 10, 75), ("الإنترنت", 15, 45), ("الهاتف", 20, 35)],
        paid_mark="✓",
    ),
}

thin = Side(style="thin", color="D1D5DB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def f(bold=False, size=11, color="111827"):
    return Font(name=FONT, bold=bold, size=size, color=color)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def header_row(ws, row, labels, col=1):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font = f(True, 11, "FFFFFF")
        c.fill = fill(TEAL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 30


def title(ws, text, sub=None, width=7):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(row=1, column=1, value=text)
    c.font = f(True, 20, "FFFFFF")
    c.fill = fill(TEAL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 42
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        s = ws.cell(row=2, column=1, value=sub)
        s.font = f(False, 11, TEAL)
        s.alignment = Alignment(horizontal="center")


def inp(c):
    c.fill = fill(YELLOW)
    c.border = BORDER
    c.font = f(color="1D4ED8")


def calc(c):
    c.border = BORDER
    c.font = f()


def build(lang, out_dir):
    t = L[lang]
    I, E = t["income"], t["expense"]
    wb = Workbook()
    names = t["sheets"]
    start, settings, tx, dash, goals, debt, bills = [wb.active] + [wb.create_sheet() for _ in range(6)]
    for ws, n in zip((start, settings, tx, dash, goals, debt, bills), names):
        ws.title = n
        ws.sheet_view.rightToLeft = t["rtl"]
        ws.sheet_view.showGridLines = False
    q = lambda n: f"'{n}'"
    S, T = q(names[1]), q(names[2])
    money = '#,##0.00;[Red]-#,##0.00;"-"'
    year = dt.date.today().year + (1 if dt.date.today().month >= 10 else 0)

    # ---- Start Here
    title(start, t["title"], t["subtitle"], 8)
    start.column_dimensions["A"].width = 110
    r = 4
    for line in t["steps"]:
        start.cell(row=r, column=1, value=line).font = f(size=12)
        r += 1
    r += 1
    c = start.cell(row=r, column=1, value=t["legend"])
    c.font = f(True, 12)
    c.fill = fill(YELLOW)
    start.cell(row=r + 1, column=1, value=t["example_note"]).font = f(size=11, color="6B7280")
    start.cell(row=r + 2, column=1, value=t["tip"]).font = f(size=11, color="6B7280")
    start.cell(row=r + 4, column=1, value="© Smart Budget Planner · personal use only").font = f(size=9, color="9CA3AF")

    # ---- Settings
    title(settings, names[1], None, 5)
    for i, (lbl, val) in enumerate([(t["currency_lbl"], "$"), (t["year_lbl"], year), (t["start_lbl"], 1000)]):
        settings.cell(row=3 + i, column=1, value=lbl).font = f(True)
        inp(settings.cell(row=3 + i, column=2, value=val))
    settings["B5"].number_format = money
    header_row(settings, 7, t["settings_hdr"])
    cat_first, cat_last = 8, 8 + len(t["cats"]) + 9  # leave 10 spare rows for user categories
    for i in range(cat_last - cat_first + 1):
        row = cat_first + i
        name, typ, budget = t["cats"][i] if i < len(t["cats"]) else ("", "", None)
        inp(settings.cell(row=row, column=1, value=name or None))
        inp(settings.cell(row=row, column=2, value={"I": I, "E": E}.get(typ)))
        b = settings.cell(row=row, column=3, value=budget if typ == "E" else None)
        inp(b)
        b.number_format = money
        d = settings.cell(row=row, column=4,
                          value=f'=IF(A{row}="","",SUMIFS({T}!$E:$E,{T}!$C:$C,A{row},{T}!$B:$B,B{row}))')
        calc(d)
        d.number_format = money
        a = settings.cell(row=row, column=5, value=f'=IF(A{row}="","",D{row}/MAX(1,MAX({T}!$F:$F)))')
        calc(a)
        a.number_format = money
    dv_type = DataValidation(type="list", formula1=f'"{I},{E}"', allow_blank=True)
    settings.add_data_validation(dv_type)
    dv_type.add(f"B{cat_first}:B{cat_last}")
    for col, w in zip("ABCDE", (26, 12, 18, 20, 16)):
        settings.column_dimensions[col].width = w
    settings.freeze_panes = "A8"

    # ---- Transactions
    title(tx, names[2], t["example_note"], 6)
    header_row(tx, 3, t["tx_hdr"])
    for i in range(ROWS):
        row = 4 + i
        for col in range(1, 6):
            inp(tx.cell(row=row, column=col))
        tx.cell(row=row, column=1).number_format = "yyyy-mm-dd"
        tx.cell(row=row, column=5).number_format = money
        m = tx.cell(row=row, column=6, value=f'=IF(A{row}="","",MONTH(A{row}))')
        calc(m)
    month_now = 1
    for i, (day, typ, cat, desc, amt) in enumerate(t["examples"]):
        row = 4 + i
        tx.cell(row=row, column=1, value=dt.date(year, month_now, day))
        tx.cell(row=row, column=2, value={"I": I, "E": E}[typ])
        tx.cell(row=row, column=3, value=cat)
        tx.cell(row=row, column=4, value=desc)
        tx.cell(row=row, column=5, value=amt)
    last = 3 + ROWS
    dv_t = DataValidation(type="list", formula1=f'"{I},{E}"', allow_blank=True)
    dv_c = DataValidation(type="list", formula1=f"={S}!$A${cat_first}:$A${cat_last}", allow_blank=True)
    dv_d = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
    for dv, rng in ((dv_t, f"B4:B{last}"), (dv_c, f"C4:C{last}"), (dv_d, f"A4:A{last}")):
        tx.add_data_validation(dv)
        dv.add(rng)
    tx.conditional_formatting.add(f"E4:E{last}", FormulaRule(formula=[f'$B4="{I}"'], font=Font(color=GREEN_TXT, bold=True)))
    for col, w in zip("ABCDEF", (14, 12, 24, 34, 14, 9)):
        tx.column_dimensions[col].width = w
    tx.freeze_panes = "A4"
    tx.auto_filter.ref = f"A3:F{last}"

    # ---- Dashboard
    title(dash, t["dash_title"], None, 12)
    dash.cell(row=3, column=2, value=t["month_pick"]).font = f(True, 12)
    mp = dash.cell(row=3, column=3, value=t["months"][0])
    inp(mp)
    mp.font = f(True, 12, "1D4ED8")
    dv_m = DataValidation(type="list", formula1='"' + ",".join(t["months"]) + '"')
    dash.add_data_validation(dv_m)
    dv_m.add("C3")
    # hidden helper: month list for MATCH
    for i, m in enumerate(t["months"]):
        dash.cell(row=40 + i, column=26, value=m)
    dash.cell(row=3, column=4, value="=MATCH(C3,$Z$40:$Z$51,0)").font = f(color="FFFFFF")
    sel = "$D$3"
    # KPI tiles row 5-6
    kpi_formulas = [
        f'=SUMIFS({T}!$E:$E,{T}!$B:$B,"{I}",{T}!$F:$F,{sel})',
        f'=SUMIFS({T}!$E:$E,{T}!$B:$B,"{E}",{T}!$F:$F,{sel})',
        "=B6-D6",
        "=IF(B6=0,0,F6/B6)",
    ]
    for i, (lbl, fm) in enumerate(zip(t["kpis"], kpi_formulas)):
        col = 2 + i * 2
        dash.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        dash.merge_cells(start_row=6, start_column=col, end_row=6, end_column=col + 1)
        h = dash.cell(row=5, column=col, value=lbl)
        h.font = f(True, 11, "FFFFFF")
        h.fill = fill(TEAL)
        h.alignment = Alignment(horizontal="center")
        v = dash.cell(row=6, column=col, value=fm)
        v.font = f(True, 18, TEAL)
        v.fill = fill(TEAL_LIGHT)
        v.alignment = Alignment(horizontal="center", vertical="center")
        v.number_format = "0.0%" if i == 3 else money
    dash.row_dimensions[6].height = 36
    # Budget vs actual (selected month)
    header_row(dash, 8, t["bva_hdr"], col=2)
    first_exp = cat_first + sum(1 for c in t["cats"] if c[1] == "I")  # skip the income categories
    n_cats = cat_last - first_exp + 1
    for i in range(n_cats):
        row, srow = 9 + i, first_exp + i
        cells = [
            f'=IF({S}!$B{srow}="{E}",{S}!$A{srow},"")',
            f'=IF(B{row}="","",{S}!$C{srow})',
            f'=IF(B{row}="","",SUMIFS({T}!$E:$E,{T}!$C:$C,B{row},{T}!$B:$B,"{E}",{T}!$F:$F,{sel}))',
            f'=IF(B{row}="","",C{row}-D{row})',
            f'=IF(B{row}="","",IF(D{row}>C{row},"{t["over"]}","{t["ok"]}"))',
        ]
        for j, fm in enumerate(cells):
            c = dash.cell(row=row, column=2 + j, value=fm)
            calc(c)
            if 1 <= j <= 3:
                c.number_format = money
    bva_last = 8 + n_cats
    dash.conditional_formatting.add(f"F9:F{bva_last}", FormulaRule(formula=["AND(ISNUMBER($D9),$D9>$C9)"], font=Font(color=RED_TXT, bold=True)))
    dash.conditional_formatting.add(f"F9:F{bva_last}", FormulaRule(formula=["AND(ISNUMBER($D9),$D9<=$C9)"], font=Font(color=GREEN_TXT, bold=True)))
    dash.conditional_formatting.add(f"E9:E{bva_last}", CellIsRule(operator="lessThan", formula=["0"], fill=fill("FEE2E2")))
    # Year overview (cols H..M)
    yr0 = 8
    header_row(dash, yr0, t["year_hdr"], col=8)
    for m in range(12):
        row = yr0 + 1 + m
        dash.cell(row=row, column=8, value=t["months"][m]).font = f(True)
        dash.cell(row=row, column=8).border = BORDER
        vals = [
            f'=SUMIFS({T}!$E:$E,{T}!$B:$B,"{I}",{T}!$F:$F,{m + 1})',
            f'=SUMIFS({T}!$E:$E,{T}!$B:$B,"{E}",{T}!$F:$F,{m + 1})',
            f"=I{row}-J{row}",
            f"=IF(I{row}=0,0,K{row}/I{row})",
            (f"={S}!$B$5+K{row}" if m == 0 else f"=M{row - 1}+K{row}"),
        ]
        for j, fm in enumerate(vals):
            c = dash.cell(row=row, column=9 + j, value=fm)
            calc(c)
            c.number_format = "0.0%" if j == 3 else money
    tot = yr0 + 13
    dash.cell(row=tot, column=8, value="Σ").font = f(True)
    for j, colL in enumerate("IJK"):
        c = dash.cell(row=tot, column=9 + j, value=f"=SUM({colL}{yr0 + 1}:{colL}{yr0 + 12})")
        c.font = f(True)
        c.number_format = money
        c.border = BORDER
    c = dash.cell(row=tot, column=12, value=f"=IF(I{tot}=0,0,K{tot}/I{tot})")
    c.font = f(True)
    c.number_format = "0.0%"
    dash.conditional_formatting.add(f"K{yr0 + 1}:K{yr0 + 12}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color=RED_TXT)))
    dash.conditional_formatting.add(f"L{yr0 + 1}:L{yr0 + 12}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=0.5, end_color="86EFAC"))
    # charts
    bar = BarChart()
    bar.title = t["chart_year"]
    bar.height, bar.width = 8, 16
    bar.add_data(Reference(dash, min_col=9, max_col=10, min_row=yr0, max_row=yr0 + 12), titles_from_data=True)
    bar.set_categories(Reference(dash, min_col=8, min_row=yr0 + 1, max_row=yr0 + 12))
    dash.add_chart(bar, f"H{bva_last + 2}")
    pie = PieChart()
    pie.title = t["chart_pie"]
    pie.height, pie.width = 8, 12
    pie_last = 8 + sum(1 for c in t["cats"] if c[1] == "E")  # default expense categories only, keeps the legend clean
    pie.add_data(Reference(dash, min_col=4, min_row=8, max_row=pie_last), titles_from_data=True)
    pie.set_categories(Reference(dash, min_col=2, min_row=9, max_row=pie_last))
    dash.add_chart(pie, f"B{bva_last + 2}")
    dash.column_dimensions["A"].width = 3
    for col, w in zip("BCDEFGHIJKLM", (24, 14, 14, 14, 16, 3, 14, 14, 14, 14, 13, 14)):
        dash.column_dimensions[col].width = w
    dash.column_dimensions["Z"].hidden = True

    # ---- Savings goals
    title(goals, names[4], None, 7)
    header_row(goals, 3, t["goals_hdr"])
    for i in range(15):
        row = 4 + i
        ex = t["goals_ex"][i] if i < len(t["goals_ex"]) else None
        for col, val in ((1, ex and ex[0]), (2, ex and ex[1]), (3, ex and ex[2]),
                         (6, ex and dt.date(year, min(12, ex[3]), 28))):
            c = goals.cell(row=row, column=col, value=val)
            inp(c)
        goals.cell(row=row, column=2).number_format = money
        goals.cell(row=row, column=3).number_format = money
        goals.cell(row=row, column=6).number_format = "yyyy-mm-dd"
        for col, fm, fmt in (
            (4, f'=IF(B{row}="","",MIN(1,C{row}/B{row}))', "0%"),
            (5, f'=IF(B{row}="","",MAX(0,B{row}-C{row}))', money),
            (7, f'=IF(OR(B{row}="",F{row}=""),"",E{row}/MAX(1,ROUNDUP((F{row}-TODAY())/30.4,0)))', money),
        ):
            c = goals.cell(row=row, column=col, value=fm)
            calc(c)
            c.number_format = fmt
    goals.conditional_formatting.add("D4:D18", ColorScaleRule(start_type="num", start_value=0, start_color="FECACA",
                                                              mid_type="num", mid_value=0.5, mid_color="FEF08A",
                                                              end_type="num", end_value=1, end_color="86EFAC"))
    for col, w in zip("ABCDEFG", (24, 15, 17, 12, 15, 16, 15)):
        goals.column_dimensions[col].width = w

    # ---- Debt payoff
    title(debt, names[5], t["debt_note"], 7)
    header_row(debt, 3, t["debt_hdr"])
    for i in range(10):
        row = 4 + i
        ex = t["debt_ex"][i] if i < len(t["debt_ex"]) else None
        for col, val, fmt in ((1, ex and ex[0], None), (2, ex and ex[1], money), (3, ex and ex[2], "0.0%"), (4, ex and ex[3], money)):
            c = debt.cell(row=row, column=col, value=val)
            inp(c)
            if fmt:
                c.number_format = fmt
        months_f = f'=IF(OR(B{row}="",D{row}=""),"",IFERROR(ROUNDUP(IF(C{row}=0,B{row}/D{row},NPER(C{row}/12,-D{row},B{row})),0),"∞"))'
        for col, fm, fmt in (
            (5, months_f, "0"),
            (6, f'=IF(ISNUMBER(E{row}),EDATE(TODAY(),E{row}),"")', "yyyy-mm"),
            (7, f'=IF(ISNUMBER(E{row}),MAX(0,E{row}*D{row}-B{row}),"")', money),
        ):
            c = debt.cell(row=row, column=col, value=fm)
            calc(c)
            c.number_format = fmt
    debt.cell(row=15, column=1, value="Σ").font = f(True)
    for col, colL in ((2, "B"), (4, "D"), (7, "G")):
        c = debt.cell(row=15, column=col, value=f"=SUM({colL}4:{colL}13)")
        c.font = f(True)
        c.number_format = money
    for col, w in zip("ABCDEFG", (22, 14, 16, 16, 18, 14, 15)):
        debt.column_dimensions[col].width = w

    # ---- Bills
    title(bills, names[6], None, 15)
    header_row(bills, 3, t["bills_hdr"] + t["months"])
    dv_p = DataValidation(type="list", formula1=f'"{t["paid_mark"]}"', allow_blank=True)
    bills.add_data_validation(dv_p)
    dv_p.add("D4:O28")
    for i in range(25):
        row = 4 + i
        ex = t["bills_ex"][i] if i < len(t["bills_ex"]) else None
        for col, val in ((1, ex and ex[0]), (2, ex and ex[1]), (3, ex and ex[2])):
            inp(bills.cell(row=row, column=col, value=val))
        bills.cell(row=row, column=3).number_format = money
        for col in range(4, 16):
            c = bills.cell(row=row, column=col)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")
            c.fill = fill(GREY)
    bills.conditional_formatting.add("D4:O28", CellIsRule(operator="equal", formula=[f'"{t["paid_mark"]}"'],
                                                          fill=fill("BBF7D0"), font=Font(color=GREEN_TXT, bold=True)))
    bills.cell(row=30, column=1, value="Σ").font = f(True)
    c = bills.cell(row=30, column=3, value="=SUM(C4:C28)")
    c.font = f(True)
    c.number_format = money
    bills.column_dimensions["A"].width = 22
    for col in range(2, 16):
        bills.column_dimensions[bills.cell(row=3, column=col).column_letter].width = 12 if col > 3 else 11

    wb.active = 0
    path = Path(out_dir) / t["file"]
    wb.save(path)
    return path


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "product")
    out.mkdir(parents=True, exist_ok=True)
    for lang in L:
        print(build(lang, out))
